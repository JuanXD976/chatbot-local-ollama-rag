from __future__ import annotations

import io
from typing import List

import openpyxl
from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse

from src.api.dependencies import build_messages_for_model, build_services
from src.api.schemas import (
    ChatRequest,
    ChatResponse,
    HealthResponse,
    MemoryResetResponse,
    RebuildResponse,
    SessionDetail,
    SessionMessage,
    SessionSummary,
)
from src.app.document_service import DocumentService
from src.attachments.chat_attachment_service import stream_answer_with_attachments
from src.config.settings import OLLAMA_CHAT_MODEL
from src.llm.ollama_client import check_ollama_connection, clean_response
from src.memory.memory_extractor import extract_memory_fact
from src.memory.memory_service import append_message_to_memory, reset_persistent_memory
from src.routing.router import detect_intent
from src.security.prompt_guard import assess_user_prompt, get_blocked_response

app = FastAPI(title="Chatbot Local API", version="6.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.get("/health", response_model=HealthResponse)
def health() -> HealthResponse:
    rag_status = DocumentService.get_rag_status()

    try:
        ollama_info = check_ollama_connection()
        models = [model["name"] for model in ollama_info.get("models", [])]

        return HealthResponse(
            ok=True,
            ollama_connected=True,
            model=OLLAMA_CHAT_MODEL,
            model_available=OLLAMA_CHAT_MODEL in models,
            document_count=rag_status["document_count"],
            indexed_chunks=rag_status["indexed_chunks"],
            vectorstore_exists=rag_status["vectorstore_exists"],
        )
    except Exception:
        return HealthResponse(
            ok=False,
            ollama_connected=False,
            model=OLLAMA_CHAT_MODEL,
            model_available=False,
            document_count=rag_status["document_count"],
            indexed_chunks=rag_status["indexed_chunks"],
            vectorstore_exists=rag_status["vectorstore_exists"],
        )


@app.post("/chat", response_model=ChatResponse)
def chat(payload: ChatRequest) -> ChatResponse:
    guard = assess_user_prompt(payload.message)
    if guard.blocked:
        blocked_answer = get_blocked_response(guard.reason)
        return ChatResponse(
            session_id=payload.session_id or "blocked_request",
            answer=blocked_answer,
            detected_intent="security_block",
            tools_used=[],
            sources=["security_layer"],
            attachments=[],
        )

    chat_service, session_service = build_services()

    session_id = chat_service.ensure_session(payload.session_id)
    messages_for_model = build_messages_for_model(
        session_service=session_service,
        session_id=session_id,
        current_prompt=payload.message,
    )

    response = chat_service.process_message(
        session_id=session_id,
        user_message=payload.message,
        messages_for_model=messages_for_model,
    )

    memory_fact = extract_memory_fact(payload.message)
    if memory_fact:
        append_message_to_memory("user", memory_fact)

    return ChatResponse(
        session_id=session_id,
        answer=clean_response(response.answer),
        detected_intent=response.detected_intent,
        tools_used=response.tools_used,
        sources=response.sources,
        attachments=[],
    )


@app.post("/chat/stream")
def chat_stream(payload: ChatRequest):
    guard = assess_user_prompt(payload.message)
    if guard.blocked:
        blocked_answer = get_blocked_response(guard.reason)

        def blocked_stream():
            yield blocked_answer

        headers = {
            "X-Session-Id": payload.session_id or "blocked_request",
            "X-Detected-Intent": "security_block",
            "X-Model": OLLAMA_CHAT_MODEL,
            "X-Sources": "security_layer",
        }
        return StreamingResponse(blocked_stream(), media_type="text/plain", headers=headers)

    chat_service, session_service = build_services()

    session_id = chat_service.ensure_session(payload.session_id)
    messages_for_model = build_messages_for_model(
        session_service=session_service,
        session_id=session_id,
        current_prompt=payload.message,
    )

    detected_intent = detect_intent(payload.message)

    def generate():
        accumulated = ""
        already_sent = ""

        for event in chat_service.stream_message(
            session_id=session_id,
            user_message=payload.message,
            messages_for_model=messages_for_model,
        ):
            chunk = event["content"]

            if event["type"] == "final":
                accumulated = chunk
            else:
                accumulated += chunk

            cleaned_full = clean_response(accumulated)

            if len(cleaned_full) > len(already_sent):
                delta = cleaned_full[len(already_sent):]
                already_sent = cleaned_full
                if delta:
                    yield delta

        memory_fact = extract_memory_fact(payload.message)
        if memory_fact:
            append_message_to_memory("user", memory_fact)

    headers = {
        "X-Session-Id": session_id,
        "X-Detected-Intent": detected_intent,
        "X-Model": OLLAMA_CHAT_MODEL,
    }
    return StreamingResponse(generate(), media_type="text/plain", headers=headers)


@app.post("/chat/attachments")
async def chat_with_attachments(
    message: str = Form(...),
    session_id: str | None = Form(default=None),
    files: List[UploadFile] = File(default=[]),
):
    guard = assess_user_prompt(message)
    if guard.blocked:
        blocked_answer = get_blocked_response(guard.reason)

        def blocked_stream():
            yield blocked_answer

        headers = {
            "X-Session-Id": session_id or "blocked_request",
            "X-Detected-Intent": "security_block",
            "X-Model": OLLAMA_CHAT_MODEL,
            "X-Attachments": "",
        }
        return StreamingResponse(blocked_stream(), media_type="text/plain", headers=headers)

    chat_service, session_service = build_services()
    ensured_session_id = chat_service.ensure_session(session_id)

    messages_for_model = build_messages_for_model(
        session_service=session_service,
        session_id=ensured_session_id,
        current_prompt=message,
    )

    file_payloads: list[tuple[str, bytes]] = []
    attachment_names: list[str] = []

    for file in files:
        content = await file.read()
        file_payloads.append((file.filename, content))
        attachment_names.append(file.filename)

    def generate():
        accumulated = ""
        already_sent = ""

        for chunk in stream_answer_with_attachments(
            user_prompt=message,
            files=file_payloads,
            messages_for_model=messages_for_model,
        ):
            accumulated += chunk
            cleaned_full = clean_response(accumulated)

            if len(cleaned_full) > len(already_sent):
                delta = cleaned_full[len(already_sent):]
                already_sent = cleaned_full
                if delta:
                    yield delta

        final_answer = clean_response(accumulated)

        # ✅ Guardar memoria si aplica
        memory_fact = extract_memory_fact(message)
        if memory_fact:
            append_message_to_memory("user", memory_fact)

        # ✅ Persistir conversación en la sesión
        session_service.append_message(
            session_id=ensured_session_id,
            role="user",
            content=message if message.strip() else "(Adjuntos enviados)",
        )
        session_service.append_message(
            session_id=ensured_session_id,
            role="assistant",
            content=final_answer,
        )

    headers = {
        "X-Session-Id": ensured_session_id,
        "X-Detected-Intent": "attachments",
        "X-Model": OLLAMA_CHAT_MODEL,
        "X-Attachment-Count": str(len(attachment_names)),
    }
    return StreamingResponse(generate(), media_type="text/plain", headers=headers)

@app.post("/exports/response")
async def export_response(
    content: str = Form(...),
    export_format: str = Form(...),
):
    normalized_format = export_format.lower().strip()

    if normalized_format == "txt":
        return StreamingResponse(
            io.BytesIO(content.encode("utf-8")),
            media_type="text/plain",
            headers={"Content-Disposition": 'attachment; filename="respuesta.txt"'},
        )

    if normalized_format == "md":
        return StreamingResponse(
            io.BytesIO(content.encode("utf-8")),
            media_type="text/markdown",
            headers={"Content-Disposition": 'attachment; filename="respuesta.md"'},
        )

    if normalized_format == "docx":
        doc = Document()
        for paragraph in content.split("\n\n"):
            doc.add_paragraph(paragraph)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": 'attachment; filename="respuesta.docx"'},
        )

    if normalized_format == "xlsx":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Respuesta"

        for idx, line in enumerate(content.splitlines(), start=1):
            ws.cell(row=idx, column=1, value=line)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="respuesta.xlsx"'},
        )

    raise HTTPException(status_code=400, detail="Formato de exportación no soportado.")


@app.get("/sessions", response_model=list[SessionSummary])
def list_sessions():
    _, session_service = build_services()
    sessions = session_service.list_sessions()

    return [
        SessionSummary(
            session_id=s.session_id,
            title=s.title,
            created_at=s.created_at,
            updated_at=s.updated_at,
        )
        for s in sessions
    ]


@app.post("/sessions", response_model=SessionSummary)
def create_session():
    _, session_service = build_services()
    session = session_service.create_session()

    return SessionSummary(
        session_id=session.session_id,
        title=session.title,
        created_at=session.created_at,
        updated_at=session.updated_at,
    )


@app.get("/sessions/{session_id}", response_model=SessionDetail)
def get_session(session_id: str):
    _, session_service = build_services()
    session = session_service.get_session(session_id)

    if not session:
        raise HTTPException(status_code=404, detail="Sesión no encontrada.")

    return SessionDetail(
        session_id=session.session_id,
        title=session.title,
        created_at=session.created_at,
        updated_at=session.updated_at,
        messages=[
            SessionMessage(
                role=m.role,
                content=m.content,
                timestamp=m.timestamp,
            )
            for m in session.messages
        ],
    )


@app.delete("/sessions/{session_id}")
def delete_session(session_id: str):
    _, session_service = build_services()
    session_service.delete_session(session_id)
    return {"ok": True}


@app.get("/documents")
def list_documents():
    return {
        "status": DocumentService.get_rag_status(),
        "documents": DocumentService.list_documents(),
    }


@app.post("/documents/upload")
async def upload_document(files: List[UploadFile] = File(...)):
    results = []

    for file in files:
        try:
            content = await file.read()
            saved_path, indexed_chunks = DocumentService.save_file_bytes(file.filename, content)
            results.append(
                {
                    "file": file.filename,
                    "saved_path": saved_path,
                    "indexed_chunks": indexed_chunks,
                }
            )
        except Exception as exc:
            results.append(
                {
                    "file": file.filename,
                    "error": str(exc),
                }
            )

    return {"results": results}


@app.delete("/documents/{filename}")
def delete_document(filename: str):
    result = DocumentService.delete_document(filename)

    if not result["deleted_file"]:
        raise HTTPException(status_code=404, detail="Documento no encontrado.")

    return {
        "ok": True,
        "deleted_chunks": result["deleted_chunks"],
    }


@app.post("/documents/rebuild", response_model=RebuildResponse)
def rebuild_documents():
    try:
        total_chunks = DocumentService.rebuild_vectorstore()
        return RebuildResponse(total_chunks=total_chunks)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc


@app.post("/memory/reset", response_model=MemoryResetResponse)
def reset_memory():
    reset_persistent_memory()
    return MemoryResetResponse(ok=True)