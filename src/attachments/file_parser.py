from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from pathlib import Path

import docx2txt
import openpyxl
from pypdf import PdfReader


@dataclass
class ParsedAttachment:
    file_name: str
    mime_group: str  # text | spreadsheet | image | unknown
    extracted_text: str
    preview_text: str


TEXT_EXTENSIONS = {".txt", ".md", ".csv"}
DOC_EXTENSIONS = {".docx", ".pdf"}
SHEET_EXTENSIONS = {".xlsx"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}


def parse_attachment(file_name: str, file_bytes: bytes) -> ParsedAttachment:
    suffix = Path(file_name).suffix.lower()

    if suffix in {".txt", ".md"}:
        text = decode_text_bytes(file_bytes)
        return ParsedAttachment(
            file_name=file_name,
            mime_group="text",
            extracted_text=text,
            preview_text=text[:1500],
        )

    if suffix == ".csv":
        text = parse_csv(file_bytes)
        return ParsedAttachment(
            file_name=file_name,
            mime_group="spreadsheet",
            extracted_text=text,
            preview_text=text[:1500],
        )

    if suffix == ".pdf":
        text = parse_pdf(file_bytes)
        return ParsedAttachment(
            file_name=file_name,
            mime_group="text",
            extracted_text=text,
            preview_text=text[:1500],
        )

    if suffix == ".docx":
        text = parse_docx(file_bytes)
        return ParsedAttachment(
            file_name=file_name,
            mime_group="text",
            extracted_text=text,
            preview_text=text[:1500],
        )

    if suffix == ".xlsx":
        text = parse_xlsx(file_bytes)
        return ParsedAttachment(
            file_name=file_name,
            mime_group="spreadsheet",
            extracted_text=text,
            preview_text=text[:1500],
        )

    if suffix in IMAGE_EXTENSIONS:
        return ParsedAttachment(
            file_name=file_name,
            mime_group="image",
            extracted_text="",
            preview_text=f"Attached visual file: {file_name}",
        )

    return ParsedAttachment(
        file_name=file_name,
        mime_group="unknown",
        extracted_text="",
        preview_text=f"Could not extract text from file {file_name}.",
    )


def decode_text_bytes(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "latin-1", "cp1252"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="ignore")


def parse_csv(file_bytes: bytes) -> str:
    text = decode_text_bytes(file_bytes)
    reader = csv.reader(io.StringIO(text))
    rows = list(reader)

    if not rows:
        return "Empty CSV."

    lines = []
    header = rows[0]
    lines.append("Detected CSV.")
    lines.append(f"Columns: {', '.join(header)}")
    lines.append("First rows:")

    for row in rows[1:11]:
        lines.append(" | ".join(row))

    return "\n".join(lines)


def parse_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    texts: list[str] = []

    for page in reader.pages:
        try:
            texts.append(page.extract_text() or "")
        except Exception:
            continue

    full_text = "\n\n".join(t.strip() for t in texts if t.strip())
    return full_text or "Could not extract text from the PDF."


def parse_docx(file_bytes: bytes) -> str:
    temp_path = Path("temp_upload_docx_parser.docx")
    temp_path.write_bytes(file_bytes)

    try:
        text = docx2txt.process(str(temp_path)) or ""
        return text.strip() or "Could not extract text from the DOCX."
    finally:
        if temp_path.exists():
            temp_path.unlink()


def parse_xlsx(file_bytes: bytes) -> str:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    lines: list[str] = []

    lines.append("Detected Excel workbook.")
    lines.append(f"Sheets: {', '.join(wb.sheetnames)}")

    for sheet_name in wb.sheetnames[:5]:
        ws = wb[sheet_name]
        lines.append(f"\nSheet: {sheet_name}")

        rows_preview = []
        for row in ws.iter_rows(min_row=1, max_row=10, values_only=True):
            normalized = ["" if value is None else str(value) for value in row]
            if any(cell != "" for cell in normalized):
                rows_preview.append(" | ".join(normalized))

        if rows_preview:
            lines.extend(rows_preview)
        else:
            lines.append("No readable content found in the first rows.")

    return "\n".join(lines)